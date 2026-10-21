"""Reference (laser diffraction) size distributions for the real test images and
the image -> reference mapping. From attempt 1 (src/real_check.py, src/alloy2_check.py)."""
from __future__ import annotations

import csv
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"

# sample name -> (image path, description, reference rows, reference source)
REAL_SAMPLES = {
    "AA7075_base": ("aa7075/Original_Image_for_Fig_2a.tif", "AA7075 base powder (Zeiss SE, 300x)",
                    ["7075"], "aa7075"),
    "AA7075_TiC_surface": ("aa7075/Original_Image_for_Fig_2c.tif", "AA7075 + TiC surface-modified (BM)",
                           ["7075 0.3wt%TiC BM", "7075 1.0wt%TiC BM", "7075 2.0wt%TiC BM"], "aa7075"),
    "AA7075_TiC_volume": ("aa7075/Original_Image_for_Fig_2d.tif", "AA7075 + TiC volume-modified (VM)",
                          ["7075 0.3wt%TiC VM", "7075 1.0wt%TiC VM"], "aa7075"),
    "Ti6Al4V": ("alloy2/Ti-6Al-4V.tif", "Ti-6Al-4V gas-atomized (TESCAN BSE)", ["Ti64"], "alloy2"),
    "Ti64_10Cu": ("alloy2/Ti64+10Cu.tif", "Ti-6Al-4V + 10 wt% Cu blend", ["Ti64+10Cu"], "alloy2"),
    "Ti64_15Cu": ("alloy2/Ti64+15Cu.tif", "Ti-6Al-4V + 15 wt% Cu blend", ["Ti64+15Cu"], "alloy2"),
    "Ti64_20Cu": ("alloy2/Ti64+20Cu.tif", "Ti-6Al-4V + 20 wt% Cu blend", ["Ti64+20Cu"], "alloy2"),
    "Cu": ("alloy2/Cu.tif", "Cu powder (TESCAN BSE)", ["Cu"], "alloy2"),
}


def load_aa7075():
    rows = list(csv.reader(open(DATA / "aa7075" / "PSD.CSV", encoding="latin-1"), delimiter=";"))
    h = rows[0]
    f = lambda x: float(x.replace(",", ".")) if x.strip() else np.nan  # noqa: E731
    out = {}
    for r in rows[2:]:
        if r and r[0].strip():
            out[r[0].strip()] = {"D10": f(r[h.index("d (0.1)")]), "D50": f(r[h.index("d (0.5)")]),
                                 "D90": f(r[h.index("d (0.9)")])}
    return out


def load_alloy2():
    import openpyxl
    wb = openpyxl.load_workbook(DATA / "alloy2" / "Table2_ParticleSizeDistribution.xlsx", data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        found = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().lower() in ("d10", "d50", "d90"):
                    v = ws.cell(row=c.row, column=c.column + 1).value
                    if isinstance(v, (int, float)):
                        found[c.value.strip().upper()] = float(v)
        out[name] = found
    return out


def reference(sample):
    _, _, refs, src = REAL_SAMPLES[sample]
    table = load_aa7075() if src == "aa7075" else load_alloy2()
    vals = [table[r] for r in refs if r in table]
    if not vals:
        return None
    return {k: float(np.mean([v[k] for v in vals if k in v])) for k in ("D10", "D50", "D90")}
