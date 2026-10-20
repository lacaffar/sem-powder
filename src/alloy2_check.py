"""Alloy-2 check: Ti-6Al-4V (and Cu, Ti64+Cu blends) from Zenodo 17037765.

Tool run unchanged on TESCAN Vega SEM images (BSE, 15 kV, real data bar).
Reference PSD: the dataset's Table2 workbook (d10/d50/d90 cells AC2:AC4 per
sheet). Measurement method is not stated in the record; the log bins starting at
0.012 µm suggest laser diffraction, so it is compared with volume-weighted stats.
Scale truth: data bar 'View field: 553.5 µm' over 1024 px = 0.5405 µm/px.

Usage: python -m src.alloy2_check [--methods watershed,cellpose]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import cv2
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
A2 = ROOT / "data/alloy2"
VIEW_FIELD_UMPX = 553.5 / 1024
SHEETS = {"Ti-6Al-4V.tif": "Ti64", "Cu.tif": "Cu", "Ti64+10Cu.tif": "Ti64+10Cu",
          "Ti64+15Cu.tif": "Ti64+15Cu", "Ti64+20Cu.tif": "Ti64+20Cu"}


def load_ref():
    import openpyxl
    wb = openpyxl.load_workbook(A2 / "Table2_ParticleSizeDistribution.xlsx", data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        found = {}
        # labels 'd10'/'d50'/'d90' with the value in the next cell to the right
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().lower() in ("d10", "d50", "d90"):
                    v = ws.cell(row=c.row, column=c.column + 1).value
                    if isinstance(v, (int, float)):
                        found[c.value.strip().upper()] = float(v)
        out[name] = {k: found.get(k, np.nan) for k in ("D10", "D50", "D90")}
    return out


def main():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from .metrics import analyze_image
    from .segment import overlay

    ap = argparse.ArgumentParser()
    ap.add_argument("--methods", default="watershed,cellpose")
    a = ap.parse_args()
    methods = a.methods.split(",")
    ref = load_ref()
    out = ROOT / "out/alloy2"
    out.mkdir(parents=True, exist_ok=True)
    rows, curves = [], {}
    for fn, sheet in SHEETS.items():
        for m in methods:
            if m != "watershed" and fn != "Ti-6Al-4V.tif":
                continue  # Cellpose only on the alloy itself (CPU time)
            df, s, lab, img, sc = analyze_image(A2 / fn, m)
            cv2.imwrite(str(out / f"{Path(fn).stem}__{m}.png"), overlay(img, lab, title=f"{fn} {m} n={s['count']}"))
            curves[(fn, m)] = df
            r = ref[sheet]
            row = {"image": fn, "method": m, "count": s["count"], "um_per_px": s["um_per_px"],
                   "scale_source": s["scale_source"], "scale_err_vs_view_field": s["um_per_px"] / VIEW_FIELD_UMPX - 1,
                   "mean_circularity": s["mean_circularity"], "irregular_frac": s["irregular_frac"]}
            for k in ("D10", "D50", "D90"):
                row[f"ref_{k}"] = r[k]
                row[f"img_{k}_num"] = s[f"{k}_num"]
                row[f"img_{k}_vol_ec"] = s[f"{k}_vol_ec"]
                row[f"err_{k}_vol_ec"] = s[f"{k}_vol_ec"] / r[k] - 1
            rows.append(row)
            print(f"{fn:15s} {m:9s} n={s['count']:4d} scale {s['um_per_px']:.4f} ({s['scale_source']}) "
                  f"vol-ec D10/50/90 {s['D10_vol_ec']:.1f}/{s['D50_vol_ec']:.1f}/{s['D90_vol_ec']:.1f} "
                  f"ref {r['D10']:.1f}/{r['D50']:.1f}/{r['D90']:.1f}", flush=True)
    res = pd.DataFrame(rows)
    res.to_csv(ROOT / "out/alloy2_vs_reported.csv", index=False)
    pd.set_option("display.width", 220)
    print(res[["image", "method", "count", "ref_D50", "img_D50_vol_ec", "err_D10_vol_ec", "err_D50_vol_ec",
               "err_D90_vol_ec", "scale_err_vs_view_field"]].round(3).to_string(index=False))

    # fig5: Ti-6Al-4V overlay + volume CDF vs reference points
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.2), gridspec_kw={"width_ratios": [1.25, 1]})
    ov = cv2.imread(str(out / "Ti-6Al-4V__watershed.png"))[..., ::-1]
    ax[0].imshow(ov)
    ax[0].axis("off")
    ax[0].set_title("Ti-6Al-4V (TESCAN, BSE), watershed, unchanged settings", fontsize=9)
    col = {"watershed": "#1f77b4", "cellpose": "#d62728"}
    for m in methods:
        if ("Ti-6Al-4V.tif", m) not in curves:
            continue
        d = curves[("Ti-6Al-4V.tif", m)]
        v = d.ecd_um.to_numpy()
        o = np.argsort(v)
        w = (v ** 3 * d.edge_weight.to_numpy())[o]
        ax[1].plot(v[o], np.cumsum(w) / w.sum(), color=col[m], lw=1.8, label=f"image, {m} (vol, edge-corr.)")
    r = ref["Ti64"]
    ax[1].plot([r["D10"], r["D50"], r["D90"]], [0.1, 0.5, 0.9], "ks", label="reported D10/D50/D90")
    ax[1].set_xlabel("diameter (µm)")
    ax[1].set_ylabel("cumulative volume fraction")
    ax[1].set_xlim(0, 70)
    ax[1].grid(alpha=0.3)
    ax[1].legend(fontsize=7, loc="lower right")
    ax[1].set_title("Ti-6Al-4V: image vs reported PSD", fontsize=9)
    fig.tight_layout()
    fig.savefig(ROOT / "figures/fig5_ti64_vs_reported.png", dpi=140)
    plt.close(fig)


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))
    main()
